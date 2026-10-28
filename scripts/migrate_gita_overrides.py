#!/usr/bin/env python3
"""One-time: build overrides/gita_insight.csv (reworked title, genre, audience) from the insights workbook,
keyed by the App Prod journey_steps id. Rows are matched by pillar + position (K1..K22, G1..G59, F1..F26)
and cross-checked against the original insight title."""
import csv
import re
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from ccl import config, sources  # noqa: E402

LETTER = {"K3": "K", "G3": "G", "F3": "F"}


def norm(s):
    return re.sub(r"[^a-z0-9]", "", (s or "").lower())


def main(xlsx="Title Revamp - with insights.xlsx"):
    config.load_env()
    steps = {r["alias"]: r for r in sources._gita_steps("en")}
    wb = openpyxl.load_workbook(config.ROOT / xlsx, read_only=True)
    out, mismatched = [], []
    for sheet, letter in LETTER.items():
        ws = wb[sheet]
        hdr = {c.value: i for i, c in enumerate(next(ws.iter_rows(min_row=1, max_row=1)))}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row[0]:
                continue
            alias = f"{letter}{int(row[0])}"
            st = steps.get(alias)
            if not st:
                sys.exit(f"{alias} not found in App Prod")
            if norm(st["original_title"]) != norm(row[hdr["Current title"]]):
                mismatched.append((alias, st["original_title"], row[hdr["Current title"]]))
            out.append({"source_id": st["id"], "alias": alias, "original_title": st["original_title"],
                        "reworked_title": row[hdr["Reworked title"]], "genre": row[hdr["Genre"]],
                        "target_audience": row[hdr["Target Audience"]]})
    config.OVERRIDES_DIR.mkdir(exist_ok=True)
    with (config.OVERRIDES_DIR / "gita_insight.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(out[0]))
        w.writeheader()
        w.writerows(out)
    print(f"wrote {len(out)} rows to overrides/gita_insight.csv")
    for m in mismatched:
        print("title differs (check):", m)


if __name__ == "__main__":
    main(*sys.argv[1:])
